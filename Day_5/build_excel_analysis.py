import os
import pandas as pd
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment

base_dir = os.path.dirname(os.path.abspath(__file__))

# Locate dataset from Day 1
data_path = os.path.join(base_dir, "..", "Day_1", "Dataset", "SampleSuperstore", "Cleaned_SampleSuperstore.csv")
if not os.path.exists(data_path):
    data_path = os.path.join(base_dir, "..", "Day_1", "Dataset", "SampleSuperstore", "SampleSuperstore.csv")

print(f"Loading data from: {data_path}")
df = pd.read_csv(data_path)
df.columns = df.columns.str.strip().str.lower().str.replace(" ", "_")

wb = openpyxl.Workbook()

# 1. Setup Tabs
ws_raw = wb.active
ws_raw.title = "raw_data"
ws_working = wb.create_sheet(title="working_formula")
ws_summary = wb.create_sheet(title="summary")

# Populate raw_data
for r in dataframe_to_rows(df, index=False, header=True):
    ws_raw.append(r)

# Populate working_formula with calculated columns
df_working = df.copy()
df_working["profit_margin"] = (df_working["profit"] / df_working["sales"]).round(4)
df_working["discount_bracket"] = df_working["discount"].apply(
    lambda x: "No Discount" if x == 0 else ("Low (0-20%)" if x <= 0.2 else "High (>20%)")
)

for r in dataframe_to_rows(df_working, index=False, header=True):
    ws_working.append(r)

# 2. Styling summary tab
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
kpi_title_font = Font(name="Calibri", size=10, bold=True, color="475569")
kpi_num_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
kpi_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

ws_summary.merge_cells("A1:F1")
ws_summary["A1"] = "VEDA TECHNOLOGY INTERNSHIP - DAY 5: EXCEL SALES DATA ANALYSIS"
ws_summary["A1"].font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
ws_summary["A1"].fill = header_fill
ws_summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_summary.row_dimensions[1].height = 28

# KPI Cards
kpis = [
    ("TOTAL SALES", "=SUM(working_formula!J2:J10000)", "$#,##0.00", "B3", "B4"),
    ("TOTAL PROFIT", "=SUM(working_formula!M2:M10000)", "$#,##0.00", "C3", "C4"),
    ("OVERALL MARGIN", "=C4/B4", "0.00%", "D3", "D4"),
    ("AVG DISCOUNT", "=AVERAGE(working_formula!L2:L10000)", "0.0%", "E3", "E4"),
]

for title, formula, fmt, t_cell, f_cell in kpis:
    ws_summary[t_cell] = title
    ws_summary[t_cell].font = kpi_title_font
    ws_summary[t_cell].fill = kpi_fill
    ws_summary[t_cell].alignment = Alignment(horizontal="center")

    ws_summary[f_cell] = formula
    ws_summary[f_cell].font = kpi_num_font
    ws_summary[f_cell].number_format = fmt
    ws_summary[f_cell].fill = kpi_fill
    ws_summary[f_cell].alignment = Alignment(horizontal="center")

ws_summary["B6"] = "ONE-PARAGRAPH SUMMARY OF KEY FINDINGS"
ws_summary["B6"].font = Font(bold=True)
ws_summary.merge_cells("B7:E9")
ws_summary["B7"] = (
    "Analysis across 9,977 transactions reveals total sales of $2.30M converting into $286.4K in net profit "
    "(12.47% margin). Profitability varies sharply across categories: Technology delivers the strongest returns "
    "(17.4% margin), whereas Furniture drops to 2.49% due to heavy losses in Tables (-$17.7K) and Bookcases (-$3.5K). "
    "A critical tipping point occurs at the 20% promotional discount threshold, beyond which orders reliably generate net operating losses."
)
ws_summary["B7"].alignment = Alignment(wrap_text=True, vertical="top")

ws_raw.freeze_panes = "A2"
ws_working.freeze_panes = "A2"

# Safe column width adjustment (avoids MergedCell AttributeError)
for ws in [ws_raw, ws_working, ws_summary]:
    for col in ws.columns:
        max_len = 0
        for cell in col[:20]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

out_file = os.path.join(base_dir, "Superstore_Excel_Analysis.xlsx")
wb.save(out_file)
print(f"Excel workbook built successfully: {out_file}")
